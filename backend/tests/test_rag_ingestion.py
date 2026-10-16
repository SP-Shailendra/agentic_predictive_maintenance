from io import BytesIO
import unittest
from unittest.mock import patch

from docx import Document
from pypdf import PdfWriter

from services.chunking_service import chunk_segments
from services.document_loader import MAX_EXTRACTED_TEXT_CHARS, DocumentLoadError, load_document
from services.rag_evaluation import score_retrieval_case, score_retrieval_suite
from services.rag_service import _reciprocal_rank_fusion, build_citation_refs, validate_citation_refs
from services.embedding_service import rerank_documents


class DocumentLoaderTests(unittest.TestCase):
    def test_markdown_preserves_headings_as_section_metadata(self) -> None:
        segments = load_document("pump.md", b"# Vibration\nInspect bearings and coupling.\n## Lubrication\nCheck oil level.\n")
        self.assertEqual(segments[0].section_title, "Vibration")
        self.assertEqual(segments[0].text, "Inspect bearings and coupling.")
        self.assertEqual(segments[1].section_title, "Lubrication")

    def test_docx_extracts_heading_paragraphs_and_table_rows(self) -> None:
        document = Document()
        document.add_heading("Bearing inspection", level=1)
        document.add_paragraph("Check for looseness.")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Signal"
        table.cell(0, 1).text = "Action"
        buffer = BytesIO()
        document.save(buffer)

        segments = load_document("pump.docx", buffer.getvalue())
        self.assertEqual(segments[0].section_title, "Bearing inspection")
        self.assertIn("Signal | Action", [segment.text for segment in segments])

    def test_pdf_without_extractable_text_fails_clearly(self) -> None:
        writer = PdfWriter()
        writer.add_blank_page(width=72, height=72)
        buffer = BytesIO()
        writer.write(buffer)
        with patch("services.document_loader._ocr_pdf_page", return_value=""):
            with self.assertRaises(DocumentLoadError) as raised:
                load_document("blank.pdf", buffer.getvalue())
        self.assertEqual(raised.exception.code, "no_extractable_text")

    def test_scanned_pdf_uses_worker_ocr_fallback(self) -> None:
        writer = PdfWriter()
        writer.add_blank_page(width=72, height=72)
        buffer = BytesIO()
        writer.write(buffer)
        with patch("services.document_loader._ocr_pdf_page", return_value="Scanned bearing inspection") as ocr:
            segments = load_document("scan.pdf", buffer.getvalue())
        self.assertEqual(segments[0].text, "Scanned bearing inspection")
        self.assertEqual(segments[0].page_number, 1)
        ocr.assert_called_once()

    def test_unsupported_or_invalid_text_is_rejected(self) -> None:
        with self.assertRaisesRegex(DocumentLoadError, "Supported formats"):
            load_document("manual.rtf", b"text")
        with self.assertRaises(DocumentLoadError) as raised:
            load_document("manual.txt", b"\xff")
        self.assertEqual(raised.exception.code, "invalid_encoding")

    def test_extracted_text_size_is_bounded(self) -> None:
        with self.assertRaises(DocumentLoadError) as raised:
            load_document("large.txt", b"x" * (MAX_EXTRACTED_TEXT_CHARS + 1))
        self.assertEqual(raised.exception.code, "extracted_text_too_large")


class ChunkingTests(unittest.TestCase):
    def test_chunking_preserves_section_and_page_with_bounded_overlap(self) -> None:
        segments = load_document("guide.md", ("# Steps\n" + "word " * 12).encode())
        chunks = chunk_segments(segments, max_words=5, overlap_words=2)
        self.assertGreater(len(chunks), 1)
        self.assertTrue(all(chunk.section_title == "Steps" for chunk in chunks))
        self.assertTrue(all(chunk.text for chunk in chunks))
        self.assertEqual([chunk.chunk_index for chunk in chunks], list(range(len(chunks))))
        self.assertEqual(len(chunks[0].text.split()[-2:]), len(chunks[1].text.split()[:2]))

    def test_chunk_parameters_are_validated(self) -> None:
        with self.assertRaises(ValueError):
            chunk_segments([], max_words=5, overlap_words=5)

    def test_adjacent_paragraphs_in_a_section_are_combined(self) -> None:
        segments = load_document("guide.txt", b"Inspect the coupling.\nCheck bearing lubrication.")
        chunks = chunk_segments(segments, max_words=20, overlap_words=2)
        self.assertEqual(len(chunks), 1)
        self.assertIn("Inspect the coupling. Check bearing lubrication.", chunks[0].text)


class RetrievalEvaluationTests(unittest.TestCase):
    def test_cross_encoder_scores_reorder_retrieval_candidates(self) -> None:
        candidates = [
            {"chunk_id": "lower", "content": "pump general maintenance"},
            {"chunk_id": "higher", "content": "bearing wear vibration checks"},
        ]

        class FakeReranker:
            def rerank(self, query, documents):
                return [0.1, 0.9]

        with patch("services.embedding_service._get_reranker", return_value=FakeReranker()):
            reranked = rerank_documents("pump bearing vibration", candidates)

        self.assertEqual([item["chunk_id"] for item in reranked], ["higher", "lower"])

    def test_reciprocal_rank_fusion_combines_and_deduplicates_candidates(self) -> None:
        vector_rows = [
            {"chunk_id": "shared", "content": "semantic match"},
            {"chunk_id": "vector-only", "content": "semantic candidate"},
        ]
        keyword_rows = [
            {"chunk_id": "shared", "content": "exact identifier"},
            {"chunk_id": "keyword-only", "content": "keyword candidate"},
        ]

        fused = _reciprocal_rank_fusion(vector_rows, keyword_rows, limit=3)

        self.assertEqual([row["chunk_id"] for row in fused], ["shared", "keyword-only", "vector-only"])
        self.assertEqual(fused[0]["content"], "semantic match")

    def test_recall_and_mrr_use_labeled_relevant_passages(self) -> None:
        case = {"relevant_phrases": ["bearing inspection", "lubrication procedure"]}
        result = score_retrieval_case(
            case["relevant_phrases"],
            ["Unrelated text", "Bearing inspection steps", "Other", "Lubrication procedure"],
            k=3,
        )
        self.assertEqual(result["recall_at_k"], 0.5)
        self.assertEqual(result["mrr_at_k"], 0.5)
        self.assertEqual(result["first_relevant_rank"], 2)

    def test_suite_reports_mean_recall_and_mrr(self) -> None:
        cases = [
            {"relevant_phrases": ["bearing"]},
            {"relevant_phrases": ["pressure"]},
        ]
        result = score_retrieval_suite(cases, [["bearing"], ["unrelated"]], k=3)
        self.assertEqual(result["case_count"], 2)
        self.assertEqual(result["mean_recall_at_k"], 0.5)
        self.assertEqual(result["mean_mrr_at_k"], 0.5)
        self.assertEqual(result["mean_ndcg_at_k"], 0.5)

    def test_ndcg_penalizes_relevant_passages_ranked_late(self) -> None:
        relevant = ["bearing procedure", "lubrication procedure"]
        early = score_retrieval_case(relevant, ["bearing procedure", "lubrication procedure"], k=2)
        late = score_retrieval_case(relevant, ["unrelated", "bearing procedure"], k=2)
        self.assertEqual(early["ndcg_at_k"], 1.0)
        self.assertLess(late["ndcg_at_k"], early["ndcg_at_k"])


class CitationValidationTests(unittest.TestCase):
    def test_citations_are_built_from_retrieved_chunks(self) -> None:
        chunks = [{"chunk_id": "chunk-1", "filename": "pump.pdf", "revision": "B", "page_number": 4}]
        citations = build_citation_refs(chunks)
        self.assertEqual(citations[0]["chunk_id"], "chunk-1")
        self.assertEqual(citations[0]["revision"], "B")
        self.assertEqual(validate_citation_refs(citations, chunks), citations)

    def test_citation_to_unretrieved_chunk_is_rejected(self) -> None:
        with self.assertRaisesRegex(ValueError, "was not retrieved"):
            validate_citation_refs(
                [{"chunk_id": "invented-chunk"}],
                [{"chunk_id": "retrieved-chunk"}],
            )


if __name__ == "__main__":
    unittest.main()